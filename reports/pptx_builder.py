from pptx import Presentation
from pptx.util import Pt
from workflows.pipeline_context import PipelineContext


class PPTXReportBuilder:
    @staticmethod
    def build(context: PipelineContext, output_path: str):
        prs = Presentation()
        profile = context.cleaned_profile
        biz = context.business_context
        story = context.storytelling_report
        kpis = context.kpi_report
        insights = context.insight_report
        causes = context.root_cause_report
        recs = context.recommendation_report

        slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = f"Executive Intelligence: {profile.filename}"
        slide.placeholders[1].text = f"{biz.domain}\n{story.executive_headline}"

        def add_bullet_slide(title: str, bullets: list[str]):
            layout = prs.slide_layouts[1]
            s = prs.slides.add_slide(layout)
            s.shapes.title.text = title
            tf = s.placeholders[1].text_frame
            tf.clear()
            for idx, b in enumerate(bullets):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                p.text = b
                p.font.size = Pt(16)

        add_bullet_slide(
            "Key Performance Indicators",
            [f"{k.name} ({k.category}): {k.formatted_value}" for k in kpis.kpis],
        )

        add_bullet_slide(
            "Executive Summary: What & Why",
            [
                f"What Happened: {story.what_happened}",
                f"Why It Happened: {story.why_it_happened}",
                f"Business Impact: {story.business_impact}",
            ],
        )

        add_bullet_slide(
            "Key Business Insights",
            [
                f"[{i.impact_level}] {i.title}: {i.business_insight}"
                for i in insights.insights[:6]
            ],
        )

        add_bullet_slide(
            "Root Cause Analysis",
            [
                f"#{rc.rank} {rc.primary_driver}: {rc.supporting_evidence}"
                for rc in causes.root_causes[:5]
            ],
        )

        add_bullet_slide(
            "Strategic & Tactical Recommendations",
            [
                f"[{r.tier}] {r.title}: {r.action} -> {r.expected_impact}"
                for r in recs.recommendations[:6]
            ],
        )

        prs.save(output_path)
